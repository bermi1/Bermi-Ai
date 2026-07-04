"""Document generation: Markdown → styled Word (.docx) output."""

import os
import re
import uuid

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from ..config import get_settings

ACCENT = RGBColor(0xB4, 0x5F, 0x2D)  # warm terracotta accent
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_CITATION_RE = re.compile(r"\s*\[\d+\]")


def _add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold** markdown spans."""
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_docx(markdown: str, strip_citations: bool = False) -> DocxDocument:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if strip_citations:
            line = _CITATION_RE.sub("", line)
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            h = doc.add_heading("", level=1)
            run = h.add_run(line[3:].strip())
            run.font.color.rgb = ACCENT
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.lstrip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, line.lstrip()[2:])
        elif re.match(r"^\s*\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\s*\d+\.\s", "", line))
        else:
            p = doc.add_paragraph()
            _add_runs(p, line)
    return doc


def save_docx(markdown: str, title: str) -> str:
    """Render markdown to .docx on disk; return the storage path."""
    settings = get_settings()
    out_dir = os.path.join(settings.upload_dir, "artifacts")
    os.makedirs(out_dir, exist_ok=True)
    safe_title = re.sub(r"[^\w\- ]", "", title).strip().replace(" ", "_") or "document"
    path = os.path.join(out_dir, f"{safe_title}_{uuid.uuid4().hex[:8]}.docx")
    markdown_to_docx(markdown).save(path)
    return path


def extract_title(markdown: str, fallback: str) -> str:
    for line in markdown.splitlines():
        if line.startswith("# "):
            return line[2:].strip()[:255]
    return fallback
